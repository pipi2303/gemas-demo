"""
import-excel.py — Import data jemaat dari Excel ke GEMAS production
Jalankan: python3 scripts/import-excel.py
"""

import sys
import json
from datetime import datetime, date

try:
    import openpyxl
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl', '-q'])
    import openpyxl

try:
    import requests
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'requests', '-q'])
    import requests

# ── Konfigurasi ───────────────────────────────────────────────────────────────
BASE_URL   = "https://gemas.raksadigital.com"
USERNAME   = "nikky"
PASSWORD   = "Nikky123"
EXCEL_FILE = "/Users/p1k3/Downloads/database_jemaat_28052026_080304.xlsx"
BATCH_SIZE = 100

# ── Helpers ───────────────────────────────────────────────────────────────────
def safe_str(v):
    if v is None: return None
    s = str(v).strip()
    return s if s and s.lower() not in ['none', 'nan', '-', ''] else None

def excel_date(v):
    if v is None: return None
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    s = safe_str(v)
    if not s: return None
    for fmt in ['%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%m/%d/%Y', '%Y/%m/%d']:
        try:
            return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except: pass
    return None

def map_gender(v):
    if not v: return 'Laki-laki'
    u = str(v).strip().upper()
    if u in ['L', 'LAKI-LAKI', 'LAKI', 'PRIA', 'M', 'MALE']: return 'Laki-laki'
    return 'Perempuan'

def map_yn(v):
    if not v: return None
    s = str(v).strip()
    if s.lower() in ['ya', 'yes', 'sudah', '1', 'true']: return 'Sudah'
    if s.lower() in ['tidak', 'no', 'belum', '0', 'false']: return 'Belum'
    return s if s else None

def calc_age(date_str):
    if not date_str: return 0
    try:
        bd = datetime.strptime(date_str, '%Y-%m-%d')
        return (datetime.now() - bd).days // 365
    except: return 0

def make_id(prefix, value):
    if not value: return None
    return f"{prefix}_{str(value).replace('.', '_').replace('-', '_').replace(' ', '_')}"

# ── Login ─────────────────────────────────────────────────────────────────────
print(f"\n🔐 Login ke {BASE_URL}...")
session = requests.Session()
r = session.post(f"{BASE_URL}/api/auth/login", json={"username": USERNAME, "password": PASSWORD})
if r.status_code != 200:
    print(f"❌ Login gagal: {r.text}")
    sys.exit(1)
print(f"✅ Login berhasil")

# ── Load sectors ──────────────────────────────────────────────────────────────
sectors_r = session.get(f"{BASE_URL}/api/data/sectors")
sectors   = sectors_r.json() if sectors_r.status_code == 200 else []
SECTOR_MAP = {s['name']: s['id'] for s in sectors}
print(f"✅ Sectors: {SECTOR_MAP}")

# ── Baca Excel ────────────────────────────────────────────────────────────────
print(f"\n📂 Membaca {EXCEL_FILE}...")
wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
ws = wb.active

headers = [str(c.value).strip() if c.value else '' for c in ws[1]]
print(f"✅ {len(headers)} kolom ditemukan")

members = []
now = datetime.now().isoformat()

for row in ws.iter_rows(min_row=2, values_only=True):
    if not any(row): continue
    r = dict(zip(headers, row))

    no_induk    = safe_str(r.get('No. Induk'))
    family_code = safe_str(r.get('Kode Keluarga'))
    sector_name = safe_str(r.get('Sektor'))
    full_name   = safe_str(r.get('Nama Lengkap')) or ''
    first_name  = safe_str(r.get('Nama Pertama')) or (full_name.split()[0] if full_name else '')
    last_name   = safe_str(r.get('Nama Belakang')) or ''
    birth_date  = excel_date(r.get('Tgl Lahir'))

    member_id  = make_id('m', no_induk) or f"m_{len(members)}_{int(datetime.now().timestamp())}"
    family_id  = make_id('fam', family_code) or f"fam_unknown_{len(members)}"
    sector_id  = SECTOR_MAP.get(sector_name, '') if sector_name else ''

    # Cari sector_id by partial match jika exact tidak ketemu
    if not sector_id and sector_name:
        for name, sid in SECTOR_MAP.items():
            if sector_name.lower() in name.lower() or name.lower() in sector_name.lower():
                sector_id = sid
                break

    members.append({
        'id':                     member_id,
        'memberNumber':           no_induk,
        'fullName':               full_name,
        'firstName':              first_name,
        'lastName':               last_name,
        'familyName':             safe_str(r.get('Nama Keluarga')),
        'degree':                 safe_str(r.get('Gelar')),
        'gender':                 map_gender(r.get('Gender')),
        'birthPlace':             safe_str(r.get('Tempat Lahir')),
        'birthDate':              birth_date or now[:10],
        'age':                    calc_age(birth_date),
        'bloodType':              safe_str(r.get('Gol. Darah')),
        'maritalStatus':          safe_str(r.get('Status Nikah')),
        'marriageDateChurch':     excel_date(r.get('Tgl Nikah Gereja')),
        'marriageDateCivil':      excel_date(r.get('Tgl Nikah Sipil')),
        'baptismStatus':          map_yn(r.get('Status Baptis')),
        'baptismPlace':           safe_str(r.get('Tempat Baptis')),
        'baptismDate':            excel_date(r.get('Tgl Baptis')),
        'sidiStatus':             map_yn(r.get('Status Sidi')),
        'sidiPlace':              safe_str(r.get('Tempat Sidi')),
        'sidiDate':               excel_date(r.get('Tgl Sidi')),
        'education':              safe_str(r.get('Pendidikan')),
        'major':                  safe_str(r.get('Jurusan')),
        'occupation':             safe_str(r.get('Pekerjaan')),
        'profession':             safe_str(r.get('Profesi')),
        'workplace':              safe_str(r.get('Tempat Kerja')),
        'languageSkills':         safe_str(r.get('Penguasaan Bahasa')),
        'skills':                 safe_str(r.get('Kompetensi/Skill')),
        'organizationExperience': safe_str(r.get('Pengalaman Organisasi')),
        'churchExperience':       safe_str(r.get('Pengalaman Gerejawi')),
        'familyRole':             safe_str(r.get('Peran Keluarga')) or 'KK',
        'position':               safe_str(r.get('Jabatan')),
        'homePhone':              safe_str(r.get('Telp Rumah')),
        'phone':                  safe_str(r.get('HP')),
        'email':                  safe_str(r.get('Email')),
        'address':                safe_str(r.get('Alamat')) or '',
        'familyCode':             family_code,
        'familyId':               family_id,
        'sectorId':               sector_id,
        'membershipType':         safe_str(r.get('Status Keanggotaan')) or 'Warga Jemaat',
        'pelkatStatus':           safe_str(r.get('Status Pelkat')),
        'membershipStatus':       safe_str(r.get('Status Aktif')) or 'Aktif',
        'notes':                  safe_str(r.get('Catatan')),
        'otherHistory':           safe_str(r.get('Riwayat Lain')),
        'joinDate':               excel_date(r.get('Tgl Bergabung')),
        'createdAt':              now,
        'updatedAt':              now,
    })

print(f"✅ {len(members)} data siap diimport\n")

# ── Import ────────────────────────────────────────────────────────────────────
total_imported  = 0
total_duplicate = 0

for i in range(0, len(members), BATCH_SIZE):
    batch = members[i:i + BATCH_SIZE]
    resp  = session.post(f"{BASE_URL}/api/admin/members/import", json={"members": batch})
    if resp.status_code == 200:
        d = resp.json()
        total_imported  += d.get('imported', 0)
        total_duplicate += d.get('duplicates', 0)
        print(f"  Batch {i//BATCH_SIZE + 1}/{-(-len(members)//BATCH_SIZE)}: ✅ imported={d.get('imported')} duplicates={d.get('duplicates')}")
    else:
        print(f"  Batch {i//BATCH_SIZE + 1} ❌ ERROR {resp.status_code}: {resp.text[:300]}")

print(f"\n{'='*50}")
print(f"✅ Import selesai!")
print(f"   Total imported  : {total_imported}")
print(f"   Total duplicate : {total_duplicate}")
print(f"{'='*50}\n")
